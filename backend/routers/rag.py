"""
Router: /api/rag
PDF & Word indexing with SSE progress + cancel support, hybrid RAG chat.
All endpoints require a valid JWT - user scope enforced for collections.
"""

import ast as _ast
import html as _html
import json
import os
import asyncio
import tempfile
import shutil
import threading
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor
from typing import AsyncGenerator, List, Optional

from sqlalchemy import delete as sa_delete, select, desc
from sqlalchemy.ext.asyncio import AsyncSession
from fastapi import APIRouter, Depends, HTTPException, Query, Request
from fastapi.responses import FileResponse, StreamingResponse
from pydantic import BaseModel, Field
from starlette.formparsers import MultiPartParser
from starlette.datastructures import Headers

from services.database import (
    ChatHistory, Collection, UploadedFile, User, get_db, AsyncSessionLocal,
)
from services.auth import get_current_user
from services.vectorstore import (
    build_vectorstore_from_pdfs,
    add_files_to_collection,
    rename_collection,
    delete_file_from_collection,
    load_vectorstore,
    load_documents,
    delete_collection,
    collection_exists,
    sanitize_collection_name,
    store_uploaded_file,
    get_uploaded_file_path,
    DEFAULT_CHUNK_SIZE,
    DEFAULT_CHUNK_OVERLAP,
)
from services.retrieval import HybridRetriever
from services.llm import (
    get_llm_from_request,
    assert_llm_ready,
    DEFAULT_MODEL, DEFAULT_OLLAMA_URL, DEFAULT_API_BASE_URL, DEFAULT_TOP_K,
    DENSE_WEIGHT, SPARSE_WEIGHT,
)
from services.rag_chain import rag_stream, clear_session, hydrate_session
from services.web_search import search

router = APIRouter()

_executor = ThreadPoolExecutor(max_workers=2)
_cancel_events: dict[str, threading.Event] = {}
MAX_UPLOAD_BYTES = 500 * 1024 * 1024

_MIME_BY_EXT = {
    ".pdf":  "application/pdf",
    ".doc":  "application/msword",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


# ── Pydantic schemas ──────────────────────────────────────────────────────────

class RagChatRequest(BaseModel):
    message:      str
    collection:   str   = "default_pdf"
    session_id:   str   = "rag_default"
    model:        str   = DEFAULT_MODEL
    ollama_url:   str   = DEFAULT_OLLAMA_URL
    temperature:  float = Field(0.3, ge=0.0, le=1.0)
    top_k:        int   = Field(default_factory=lambda: DEFAULT_TOP_K, ge=1, le=50)
    use_web:      bool  = False
    provider:     str   = "ollama"
    api_base_url: str   = DEFAULT_API_BASE_URL
    api_token:    str   = ""


class RenameCollectionRequest(BaseModel):
    display_name: str


# ── DB helpers ────────────────────────────────────────────────────────────────

async def _get_collection_row(db: AsyncSession, user_id: int, folder_id: str) -> Collection:
    result = await db.execute(
        select(Collection).where(
            Collection.user_id == user_id,
            Collection.folder_id == folder_id,
        )
    )
    row = result.scalar_one_or_none()
    if not row:
        raise HTTPException(404, detail=f"Collection '{folder_id}' not found.")
    return row


async def _ensure_collection_row(
    db: AsyncSession, user_id: int, folder_id: str, display_name: str,
) -> Collection:
    result = await db.execute(
        select(Collection).where(
            Collection.user_id == user_id,
            Collection.folder_id == folder_id,
        )
    )
    row = result.scalar_one_or_none()
    if row:
        return row
    row = Collection(user_id=user_id, folder_id=folder_id, display_name=display_name)
    db.add(row)
    await db.flush()
    await db.refresh(row)
    return row


# ── Collections ───────────────────────────────────────────────────────────────

@router.get("/collections")
async def get_collections(
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    result = await db.execute(
        select(Collection)
        .where(Collection.user_id == current_user.id)
        .order_by(Collection.created_at)
    )
    rows = result.scalars().all()
    return {
        "collections": [
            {"id": r.folder_id, "display_name": r.display_name}
            for r in rows
        ]
    }


@router.get("/collections/{collection}/documents")
async def get_collection_documents(
    collection: str,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    safe_col = sanitize_collection_name(collection)
    await _get_collection_row(db, current_user.id, safe_col)

    docs = load_documents(safe_col, user_id=current_user.id)
    if not docs:
        return {"collection": safe_col, "documents": []}

    groups: dict[str, list] = defaultdict(list)
    for doc in docs:
        fname = doc.metadata.get("source_file") or doc.metadata.get("source") or "unknown"
        fname = os.path.basename(fname)
        groups[fname].append(doc)

    result_list = []
    for fname, chunks in groups.items():
        ext = os.path.splitext(fname)[1].lower()
        file_type = "pdf" if ext == ".pdf" else "word" if ext in (".docx", ".doc") else "other"
        pages = [c.metadata.get("page") for c in chunks if isinstance(c.metadata.get("page"), int)]
        page_count = (max(pages) + 1) if pages else None
        preview = chunks[0].page_content[:300].strip() if chunks else ""
        result_list.append({
            "filename": fname,
            "file_type": file_type,
            "chunk_count": len(chunks),
            "page_count": page_count,
            "preview": preview,
        })

    result_list.sort(key=lambda x: (0 if x["file_type"] == "pdf" else 1, x["filename"].lower()))
    return {"collection": safe_col, "document_count": len(result_list), "documents": result_list}


@router.delete("/collections/{collection}")
async def remove_collection(
    collection: str,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    safe_col = sanitize_collection_name(collection)
    row = await _get_collection_row(db, current_user.id, safe_col)
    delete_collection(safe_col, user_id=current_user.id)
    await db.delete(row)
    return {"status": "deleted", "collection": collection}


@router.patch("/collections/{collection}/rename")
async def rename_collection_endpoint(
    collection: str,
    body: RenameCollectionRequest,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    safe_col = sanitize_collection_name(collection)
    row = await _get_collection_row(db, current_user.id, safe_col)
    name = body.display_name.strip()
    rename_collection(safe_col, name, user_id=current_user.id)
    row.display_name = name
    db.add(row)
    return {"status": "renamed", "collection": collection, "display_name": name}


@router.delete("/collections/{collection}/files/{filename:path}")
async def remove_file_from_collection(
    collection: str,
    filename: str,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    safe_col = sanitize_collection_name(collection)
    await _get_collection_row(db, current_user.id, safe_col)
    try:
        remaining = delete_file_from_collection(safe_col, filename, user_id=current_user.id)
    except FileNotFoundError as e:
        raise HTTPException(404, detail=str(e))

    await db.execute(
        sa_delete(UploadedFile).where(
            UploadedFile.user_id == current_user.id,
            UploadedFile.filename == filename,
        )
    )

    if remaining == 0:
        try:
            row = await _get_collection_row(db, current_user.id, safe_col)
            await db.delete(row)
        except HTTPException:
            pass

    return {"status": "deleted", "filename": filename, "remaining_chunks": remaining}


@router.get("/collections/{collection}/files/{filename:path}/download")
async def download_file(
    collection: str,
    filename: str,
    inline: bool = Query(False, description="Open inline in browser when possible"),
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    safe_col = sanitize_collection_name(collection)
    await _get_collection_row(db, current_user.id, safe_col)
    path = get_uploaded_file_path(current_user.id, safe_col, filename)
    if not path:
        raise HTTPException(404, detail="Original file not found.")

    ext = os.path.splitext(filename)[1].lower()
    media_type = _MIME_BY_EXT.get(ext, "application/octet-stream")
    return FileResponse(
        path,
        media_type=media_type,
        filename=filename,
        content_disposition_type="inline" if inline else "attachment",
    )


def _word_file_to_html(path: str) -> str:
    """Convert Word document to simple readable HTML for in-app preview."""
    lower = path.lower()
    parts: list[str] = []

    if lower.endswith(".docx"):
        try:
            from docx import Document
            doc = Document(path)
            for p in doc.paragraphs:
                text = (p.text or "").strip()
                if text:
                    parts.append(f"<p>{_html.escape(text)}</p>")
            for table in doc.tables:
                rows_html = []
                for row in table.rows:
                    cells = "".join(
                        f"<td>{_html.escape((c.text or '').strip())}</td>"
                        for c in row.cells
                    )
                    rows_html.append(f"<tr>{cells}</tr>")
                if rows_html:
                    parts.append(
                        "<table class='word-preview-table'>"
                        + "".join(rows_html)
                        + "</table>"
                    )
        except Exception:
            parts = []

    if not parts:
        try:
            import docx2txt
            text = docx2txt.process(path) or ""
            for line in text.splitlines():
                line = line.strip()
                if line:
                    parts.append(f"<p>{_html.escape(line)}</p>")
        except Exception as e:
            raise HTTPException(500, detail=f"خطا در خواندن فایل Word: {e}")

    return "".join(parts) or "<p>(محتوایی برای نمایش یافت نشد)</p>"


@router.get("/collections/{collection}/files/{filename:path}/preview")
async def preview_file(
    collection: str,
    filename: str,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """
    In-app preview metadata.
    PDF: client fetches download blob into an iframe.
    Word: returns HTML rendered from document text.
    """
    safe_col = sanitize_collection_name(collection)
    await _get_collection_row(db, current_user.id, safe_col)
    path = get_uploaded_file_path(current_user.id, safe_col, filename)
    if not path:
        raise HTTPException(404, detail="Original file not found.")

    ext = os.path.splitext(filename)[1].lower()
    if ext == ".pdf":
        return {"type": "pdf", "filename": filename}
    if ext in (".doc", ".docx"):
        html_body = await asyncio.to_thread(_word_file_to_html, path)
        return {"type": "word", "filename": filename, "html": html_body}

    raise HTTPException(400, detail="پیش‌نمایش این نوع فایل پشتیبانی نمی‌شود.")


# ── Multipart helper ──────────────────────────────────────────────────────────

async def _parse_multipart(request: Request):
    headers = Headers(scope={"type": "http", "headers": request.scope["headers"]})
    parser = MultiPartParser(
        headers=headers, stream=request.stream(),
        max_files=100, max_fields=20, max_part_size=MAX_UPLOAD_BYTES,
    )
    form_data = await parser.parse()
    tmp_dir = tempfile.mkdtemp()
    saved: list[str] = []
    try:
        collection = form_data.get("collection", "default_pdf")
        chunk_size_s = form_data.get("chunk_size", str(DEFAULT_CHUNK_SIZE))
        chunk_overlap_s = form_data.get("chunk_overlap", str(DEFAULT_CHUNK_OVERLAP))
        job_id = form_data.get("job_id", "")
        for key, val in form_data.multi_items():
            if key == "files" and hasattr(val, "filename") and \
               val.filename.lower().endswith((".pdf", ".docx", ".doc")):
                dest = os.path.join(tmp_dir, val.filename)
                with open(dest, "wb") as f:
                    f.write(await val.read())
                saved.append(dest)
    finally:
        await form_data.close()
    return tmp_dir, saved, collection, chunk_size_s, chunk_overlap_s, job_id


# ── Persist uploaded files ────────────────────────────────────────────────────

async def _persist_files_async(
    prepared: list[tuple[str, str, int, str]],
    user_id: int,
    safe_col: str,
    collection: str,
) -> None:
    """Insert/update UploadedFile rows. Must run on the app event loop."""
    async with AsyncSessionLocal() as sess:
        result = await sess.execute(
            select(Collection).where(
                Collection.user_id == user_id,
                Collection.folder_id == safe_col,
            )
        )
        row = result.scalar_one_or_none()
        if not row:
            row = Collection(user_id=user_id, folder_id=safe_col, display_name=collection)
            sess.add(row)
            await sess.flush()

        for fname, ftype, fsize, dest in prepared:
            await sess.execute(
                sa_delete(UploadedFile).where(
                    UploadedFile.collection_id == row.id,
                    UploadedFile.filename == fname,
                )
            )
            sess.add(UploadedFile(
                collection_id=row.id,
                user_id=user_id,
                filename=fname,
                file_type=ftype,
                file_size=fsize,
                storage_path=dest,
            ))
        await sess.commit()


def _persist_files_from_thread(
    loop: asyncio.AbstractEventLoop,
    saved_paths: list[str],
    user_id: int,
    safe_col: str,
    collection: str,
) -> None:
    """
    Copy files on the worker thread, then schedule DB writes on the app loop.
    Avoids asyncio.run() which breaks asyncpg (engine bound to another loop).
    """
    prepared: list[tuple[str, str, int, str]] = []
    for path in saved_paths:
        fname = os.path.basename(path)
        ext = os.path.splitext(fname)[1].lower()
        ftype = "pdf" if ext == ".pdf" else "word"
        fsize = os.path.getsize(path) if os.path.exists(path) else 0
        dest = store_uploaded_file(path, fname, user_id, safe_col)
        prepared.append((fname, ftype, fsize, dest))

    fut = asyncio.run_coroutine_threadsafe(
        _persist_files_async(prepared, user_id, safe_col, collection),
        loop,
    )
    fut.result(timeout=120)


# ── Index stream ──────────────────────────────────────────────────────────────

@router.post("/index/stream")
async def index_pdfs_stream(
    request: Request,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    try:
        tmp_dir, saved_paths, collection, chunk_size_s, chunk_overlap_s, job_id = \
            await _parse_multipart(request)
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(400, detail=f"Error receiving files: {e}")

    if not saved_paths:
        shutil.rmtree(tmp_dir, ignore_errors=True)
        raise HTTPException(400, detail="No valid PDF or Word file received.")

    chunk_size = int(chunk_size_s) if str(chunk_size_s).isdigit() else DEFAULT_CHUNK_SIZE
    chunk_overlap = int(chunk_overlap_s) if str(chunk_overlap_s).isdigit() else DEFAULT_CHUNK_OVERLAP
    safe_col = sanitize_collection_name(collection)
    user_id = current_user.id
    cancel_event = threading.Event()
    if job_id:
        _cancel_events[job_id] = cancel_event

    # Pre-register collection row
    await _ensure_collection_row(db, user_id, safe_col, collection)

    queue: asyncio.Queue = asyncio.Queue()
    loop = asyncio.get_event_loop()

    def progress_cb(pct: int, detail: str):
        loop.call_soon_threadsafe(queue.put_nowait, {"progress": pct, "detail": detail})

    def run_indexing():
        try:
            _, total_chunks = build_vectorstore_from_pdfs(
                saved_paths, safe_col, chunk_size, chunk_overlap,
                display_name=collection,
                progress_callback=progress_cb,
                cancel_event=cancel_event,
                user_id=user_id,
            )
            _persist_files_from_thread(loop, saved_paths, user_id, safe_col, collection)

            if cancel_event.is_set():
                loop.call_soon_threadsafe(queue.put_nowait, {"cancelled": True})
            else:
                loop.call_soon_threadsafe(queue.put_nowait, {
                    "done": True, "total_chunks": total_chunks,
                    "files": len(saved_paths), "collection": safe_col,
                    "original_name": collection,
                })
        except Exception as e:
            loop.call_soon_threadsafe(queue.put_nowait, {"error": str(e)})
        finally:
            shutil.rmtree(tmp_dir, ignore_errors=True)
            _cancel_events.pop(job_id, None)

    _executor.submit(run_indexing)

    async def generate() -> AsyncGenerator[str, None]:
        try:
            while True:
                try:
                    event = await asyncio.wait_for(queue.get(), timeout=300.0)
                except asyncio.TimeoutError:
                    yield f"data: {json.dumps({'error': 'timeout'})}\n\n"
                    cancel_event.set()
                    break
                yield f"data: {json.dumps(event, ensure_ascii=False)}\n\n"
                if "done" in event or "error" in event or "cancelled" in event:
                    break
        except asyncio.CancelledError:
            cancel_event.set()

    return StreamingResponse(generate(), media_type="text/event-stream")


@router.delete("/index/{job_id}")
def cancel_index(job_id: str, _: User = Depends(get_current_user)):
    ev = _cancel_events.get(job_id)
    if ev:
        ev.set()
        return {"status": "cancellation_requested"}
    return {"status": "not_found"}


# ── Add files to existing collection ─────────────────────────────────────────

@router.post("/collections/{collection}/files/stream")
async def add_files_to_collection_stream(
    collection: str,
    request: Request,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    safe_col = sanitize_collection_name(collection)
    await _get_collection_row(db, current_user.id, safe_col)

    try:
        _, saved_paths, _, chunk_size_s, chunk_overlap_s, job_id = \
            await _parse_multipart(request)
    except Exception as e:
        raise HTTPException(400, detail=f"Error receiving files: {e}")

    if not saved_paths:
        raise HTTPException(400, detail="No valid file received.")

    chunk_size = int(chunk_size_s) if str(chunk_size_s).isdigit() else DEFAULT_CHUNK_SIZE
    chunk_overlap = int(chunk_overlap_s) if str(chunk_overlap_s).isdigit() else DEFAULT_CHUNK_OVERLAP
    user_id = current_user.id
    cancel_event = threading.Event()
    if job_id:
        _cancel_events[job_id] = cancel_event

    queue: asyncio.Queue = asyncio.Queue()
    loop = asyncio.get_event_loop()

    def progress_cb(pct: int, detail: str):
        loop.call_soon_threadsafe(queue.put_nowait, {"progress": pct, "detail": detail})

    def run_adding():
        tmp_dir = os.path.dirname(saved_paths[0])
        try:
            _, total_chunks = add_files_to_collection(
                saved_paths, safe_col, chunk_size, chunk_overlap,
                progress_callback=progress_cb,
                cancel_event=cancel_event,
                user_id=user_id,
            )
            _persist_files_from_thread(loop, saved_paths, user_id, safe_col, collection)

            if cancel_event.is_set():
                loop.call_soon_threadsafe(queue.put_nowait, {"cancelled": True})
            else:
                loop.call_soon_threadsafe(queue.put_nowait, {
                    "done": True, "total_chunks": total_chunks,
                    "files": len(saved_paths), "collection": safe_col,
                })
        except Exception as e:
            loop.call_soon_threadsafe(queue.put_nowait, {"error": str(e)})
        finally:
            shutil.rmtree(tmp_dir, ignore_errors=True)
            _cancel_events.pop(job_id, None)

    _executor.submit(run_adding)

    async def generate() -> AsyncGenerator[str, None]:
        try:
            while True:
                try:
                    event = await asyncio.wait_for(queue.get(), timeout=300.0)
                except asyncio.TimeoutError:
                    yield f"data: {json.dumps({'error': 'timeout'})}\n\n"
                    cancel_event.set()
                    break
                yield f"data: {json.dumps(event, ensure_ascii=False)}\n\n"
                if "done" in event or "error" in event or "cancelled" in event:
                    break
        except asyncio.CancelledError:
            cancel_event.set()

    return StreamingResponse(generate(), media_type="text/event-stream")


# ── RAG Chat ──────────────────────────────────────────────────────────────────

@router.post("/chat/stream")
async def rag_chat_stream(
    req: RagChatRequest,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    safe_col = sanitize_collection_name(req.collection)
    col_row = await _get_collection_row(db, current_user.id, safe_col)

    if not collection_exists(safe_col, user_id=current_user.id):
        raise HTTPException(404, detail=f"Collection '{req.collection}' not indexed yet.")
    assert_llm_ready(
        provider=req.provider,
        ollama_url=req.ollama_url,
        api_base_url=req.api_base_url,
        api_token=req.api_token,
    )

    user_id = current_user.id
    col_id = col_row.id
    session_key = req.session_id

    async def generate() -> AsyncGenerator[str, None]:
        try:
            vs = load_vectorstore(safe_col, user_id=user_id)
            docs = load_documents(safe_col, user_id=user_id)
            retriever = HybridRetriever(
                vectorstore=vs, documents=docs,
                dense_k=max(req.top_k * 2, 16),
                sparse_k=max(req.top_k * 2, 16),
                final_k=req.top_k,
                dense_weight=DENSE_WEIGHT,
                sparse_weight=SPARSE_WEIGHT,
            )
            llm = get_llm_from_request(req)
            user_input = req.message
            web_results: list = []

            if req.use_web:
                yield f"data: {json.dumps({'status': 'searching', 'msg': 'Searching the web...'})}\n\n"
                await asyncio.sleep(0)
                web_results = await asyncio.to_thread(search, req.message, 10)
                if web_results:
                    yield f"data: {json.dumps({'status': 'search_done', 'msg': f'{len(web_results)} results found'})}\n\n"
                    user_input += "\n\n[Web Search Results]\n" + "\n".join(
                        r["body"] for r in web_results if r.get("body")
                    )
                else:
                    yield f"data: {json.dumps({'status': 'search_done', 'msg': 'No results found'})}\n\n"
                await asyncio.sleep(0)

            sources_data = None
            full_answer = ""

            async for token in rag_stream(
                llm=llm, retriever=retriever,
                question=user_input, session_id=req.session_id,
                executor=_executor,
            ):
                if token.startswith("\x00STATUS\x00"):
                    payload = token[len("\x00STATUS\x00"):]
                    if payload == "searching":
                        yield f"data: {json.dumps({'status': 'retrieving', 'msg': 'Searching knowledge base...'})}\n\n"
                    elif payload.startswith("done|"):
                        count = payload.split("|")[1]
                        yield f"data: {json.dumps({'status': 'retrieval_done', 'msg': f'{count} sources found'})}\n\n"
                    await asyncio.sleep(0)
                    continue

                if token.startswith("\x00SOURCES\x00"):
                    try:
                        sources_data = _ast.literal_eval(token[len("\x00SOURCES\x00"):])
                    except Exception:
                        pass
                    continue

                full_answer += token
                yield f"data: {json.dumps({'token': token}, ensure_ascii=False)}\n\n"
                await asyncio.sleep(0)

            # Persist chat history
            sources_json = json.dumps(sources_data, ensure_ascii=False) if sources_data else None
            async with AsyncSessionLocal() as hist_db:
                hist_db.add(ChatHistory(
                    user_id=user_id, collection_id=col_id, session_key=session_key,
                    role="user", content=req.message,
                ))
                hist_db.add(ChatHistory(
                    user_id=user_id, collection_id=col_id, session_key=session_key,
                    role="assistant", content=full_answer, sources=sources_json,
                ))
                await hist_db.commit()

            if sources_data:
                yield f"data: {json.dumps({'sources': sources_data}, ensure_ascii=False)}\n\n"

            if web_results:
                web_sources = [
                    {"title": r.get("title", ""), "body": r.get("body", "")[:200], "link": r.get("link", "")}
                    for r in web_results if r.get("body") or r.get("link")
                ]
                yield f"data: {json.dumps({'web_sources': web_sources}, ensure_ascii=False)}\n\n"

            yield f"data: {json.dumps({'done': True})}\n\n"

        except asyncio.CancelledError:
            pass
        except Exception as e:
            yield f"data: {json.dumps({'error': str(e)})}\n\n"

    return StreamingResponse(generate(), media_type="text/event-stream")


@router.post("/clear")
async def clear_rag_session(
    session_id: str = "rag_default",
    _: User = Depends(get_current_user),
):
    clear_session(session_id)
    return {"status": "cleared"}


# ── Chat History ──────────────────────────────────────────────────────────────

@router.get("/collections/{collection}/history/sessions")
async def list_chat_sessions(
    collection: str,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """List distinct conversations for a collection (ChatGPT-style thread list)."""
    safe_col = sanitize_collection_name(collection)
    col_row = await _get_collection_row(db, current_user.id, safe_col)

    result = await db.execute(
        select(ChatHistory)
        .where(
            ChatHistory.user_id == current_user.id,
            ChatHistory.collection_id == col_row.id,
        )
        .order_by(ChatHistory.created_at.asc())
    )
    rows = result.scalars().all()

    sessions: dict[str, dict] = {}
    for r in rows:
        sk = r.session_key
        if sk not in sessions:
            sessions[sk] = {
                "session_key": sk,
                "title": None,
                "created_at": r.created_at,
                "updated_at": r.created_at,
                "message_count": 0,
            }
        entry = sessions[sk]
        entry["message_count"] += 1
        entry["updated_at"] = r.created_at
        if entry["title"] is None and r.role == "user" and r.content.strip():
            text = r.content.strip().replace("\n", " ")
            entry["title"] = (text[:80] + "…") if len(text) > 80 else text

    ordered = sorted(sessions.values(), key=lambda s: s["updated_at"], reverse=True)
    return {
        "collection": safe_col,
        "sessions": [
            {
                "session_key": s["session_key"],
                "title": s["title"] or "گفتگوی بدون عنوان",
                "created_at": s["created_at"].isoformat(),
                "updated_at": s["updated_at"].isoformat(),
                "message_count": s["message_count"],
            }
            for s in ordered
        ],
    }


@router.get("/collections/{collection}/history")
async def get_chat_history(
    collection: str,
    limit: int = 50,
    session: Optional[str] = Query(None, description="Filter to one session_key"),
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    safe_col = sanitize_collection_name(collection)
    col_row = await _get_collection_row(db, current_user.id, safe_col)

    q = select(ChatHistory).where(
        ChatHistory.user_id == current_user.id,
        ChatHistory.collection_id == col_row.id,
    )
    if session:
        q = q.where(ChatHistory.session_key == session)

    if session:
        # Chronological for one conversation (id breaks timestamp ties)
        result = await db.execute(
            q.order_by(ChatHistory.created_at.asc(), ChatHistory.id.asc()).limit(limit * 2)
        )
        rows = list(result.scalars().all())
    else:
        # Latest messages across sessions, then chronological within the page
        result = await db.execute(
            q.order_by(desc(ChatHistory.created_at), desc(ChatHistory.id)).limit(limit * 2)
        )
        rows = list(reversed(result.scalars().all()))
    return {
        "collection": safe_col,
        "session": session,
        "history": [
            {
                "id": r.id,
                "role": r.role,
                "content": r.content,
                "sources": json.loads(r.sources) if r.sources else None,
                "session": r.session_key,
                "created_at": r.created_at.isoformat(),
            }
            for r in rows
        ],
    }


@router.delete("/collections/{collection}/history/sessions/{session_key}")
async def delete_chat_session(
    collection: str,
    session_key: str,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    safe_col = sanitize_collection_name(collection)
    col_row = await _get_collection_row(db, current_user.id, safe_col)
    await db.execute(
        sa_delete(ChatHistory).where(
            ChatHistory.user_id == current_user.id,
            ChatHistory.collection_id == col_row.id,
            ChatHistory.session_key == session_key,
        )
    )
    clear_session(session_key)
    return {"status": "deleted", "session_key": session_key}


@router.post("/collections/{collection}/history/sessions/{session_key}/hydrate")
async def hydrate_chat_session(
    collection: str,
    session_key: str,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Load DB messages into in-memory LangChain history for resumed chats."""
    safe_col = sanitize_collection_name(collection)
    col_row = await _get_collection_row(db, current_user.id, safe_col)

    result = await db.execute(
        select(ChatHistory)
        .where(
            ChatHistory.user_id == current_user.id,
            ChatHistory.collection_id == col_row.id,
            ChatHistory.session_key == session_key,
        )
        .order_by(ChatHistory.created_at.asc(), ChatHistory.id.asc())
    )
    rows = result.scalars().all()
    if not rows:
        raise HTTPException(404, detail="گفتگویی با این شناسه یافت نشد.")

    hydrate_session(
        session_key,
        [{"role": r.role, "content": r.content} for r in rows],
    )
    return {"status": "hydrated", "session_key": session_key, "messages": len(rows)}


@router.delete("/collections/{collection}/history")
async def clear_chat_history(
    collection: str,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    safe_col = sanitize_collection_name(collection)
    col_row = await _get_collection_row(db, current_user.id, safe_col)

    result = await db.execute(
        select(ChatHistory.session_key)
        .where(
            ChatHistory.user_id == current_user.id,
            ChatHistory.collection_id == col_row.id,
        )
        .distinct()
    )
    for (sk,) in result.all():
        clear_session(sk)

    await db.execute(
        sa_delete(ChatHistory).where(
            ChatHistory.user_id == current_user.id,
            ChatHistory.collection_id == col_row.id,
        )
    )
    return {"status": "cleared", "collection": safe_col}
